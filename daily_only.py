"""Retire the independent monthly ledger in the served app.

Only daily_profit and extra_profit are authoritative. monthly_summary remains an
inert historical archive until a verified, independently backed-up migration.
Do not drop, rewrite, or silently backfill it on a production startup.
"""
import calendar
import datetime as dt
import io
import re
from functools import wraps

from flask import flash, redirect, request, session, url_for

from daily_monthly_ledger import install as install_derived_months


DAY_NAMES = '월화수목금토일'
MAX_WORKBOOK_SIZE = 10 * 1024 * 1024


def _amount(value):
    """Do not turn malformed financial cells into fabricated zeroes."""
    if value is None or value == '':
        return 0
    if isinstance(value, bool):
        raise ValueError('금액 셀에 숫자를 입력해 주세요.')
    text = str(value).strip().replace(',', '')
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError('금액은 정수여야 합니다.')
        return int(value)
    if not re.fullmatch(r'[+-]?\d+', text):
        raise ValueError('금액 셀에 정수만 입력해 주세요.')
    return int(text)


def _daily_upsert(conn, user_id, date, disp, daily, nim, total, dpd=None, memo=None):
    """Portable SQLite/PostgreSQL upsert preserving original IDs and weekly memos."""
    columns = ['user_id', 'date', 'day_of_week', 'dispensing_fee',
               'daily_net_profit', 'dispensing_plus_daily', 'non_insurance_margin',
               'total', 'updated_at']
    row = [user_id, date.isoformat(), DAY_NAMES[date.weekday()], disp, daily,
           disp + daily if dpd is None else dpd, nim, total]
    updates = ', '.join(f'{name}=EXCLUDED.{name}' for name in columns[2:-1])
    if memo is not None:
        columns.append('memo')
        row.append(memo)
        updates += ', memo=EXCLUDED.memo'
    placeholders = ', '.join('?' for _ in row) + ', CURRENT_TIMESTAMP'
    # Updated_at appears before memo in the column list. Keep parameter order
    # aligned by building the VALUES list independently of optional memo.
    if memo is not None:
        placeholders = ', '.join('?' for _ in row[:-1]) + ', CURRENT_TIMESTAMP, ?'
    sql = (f"INSERT INTO daily_profit ({', '.join(columns)}) VALUES ({placeholders}) "
           f'ON CONFLICT (user_id, date) DO UPDATE SET {updates}, updated_at=CURRENT_TIMESTAMP')
    conn.execute(sql, row)


def _open_workbook(file_bytes, password):
    import openpyxl
    if file_bytes.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        if not password:
            raise ValueError('암호가 있는 엑셀 파일은 암호를 입력해 주세요.')
        import msoffcrypto
        decrypted = io.BytesIO()
        office = msoffcrypto.OfficeFile(io.BytesIO(file_bytes))
        office.load_key(password=password)
        office.decrypt(decrypted)
        decrypted.seek(0)
        return openpyxl.load_workbook(decrypted, data_only=True, read_only=True)
    return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)


def _parse_daily_rows(workbook, filename, start_year):
    """Yield daily cells only. Never inspect/import any month-summary worksheet."""
    ws = workbook.active
    headers = [str(ws.cell(1, c).value or '') for c in range(1, 7)]
    standard = any('날짜' in h for h in headers) and any('조제' in h for h in headers)
    if standard:
        for cells in ws.iter_rows(min_row=2, max_col=6, values_only=True):
            raw_date = cells[0]
            if raw_date is None or str(raw_date).strip() == '':
                continue
            try:
                if isinstance(raw_date, dt.datetime):
                    day = raw_date.date()
                elif isinstance(raw_date, dt.date):
                    day = raw_date
                else:
                    day = dt.date.fromisoformat(str(raw_date).strip()[:10])
            except ValueError as exc:
                raise ValueError(f'올바르지 않은 일별 날짜: {raw_date}') from exc
            disp, daily, nim = (_amount(cells[i]) for i in (2, 3, 4))
            yield day, disp, daily, nim, disp + daily + nim, None, str(cells[5] or '')
        return

    # Old weekly-only sheets do not identify the year in their week header.
    # Infer it from explicit form input or filename; NEVER from the retired
    # second-sheet month ledger or a guessed year.
    year_text = str(start_year or '').strip()
    if not year_text:
        match = re.search(r'(20\d{2})', filename or '')
        year_text = match.group(1) if match else ''
    if not re.fullmatch(r'20\d{2}', year_text) or not 2000 <= int(year_text) <= 2100:
        raise ValueError('기존 주표는 파일명에 연도(예: 2026)를 넣거나 시작 연도를 지정해 주세요.')
    year, month, start_day, offset = int(year_text), 10, 17, 0
    found_header = False
    for cells in ws.iter_rows(max_col=6, values_only=True):
        first = cells[0]
        if first is None:
            offset = 0
            continue
        label = str(first).strip()
        header = re.search(r'(\d{1,2})\s*월\s*(\d{1,2})\s*일?\s*~', label)
        if header:
            next_month, next_day = int(header.group(1)), int(header.group(2))
            if not 1 <= next_month <= 12 or not 1 <= next_day <= 31:
                raise ValueError('주표의 월·일이 올바르지 않습니다.')
            if next_month < month and next_month <= 2 and month >= 11:
                year += 1
            month, start_day, offset = next_month, next_day, 0
            found_header = True
            continue
        if label not in DAY_NAMES:
            continue
        if not found_header:
            raise ValueError('주표의 첫 주 시작 날짜를 확인해 주세요.')
        disp, daily = _amount(cells[1]), _amount(cells[2])
        combined, nim, total = (_amount(cells[i]) for i in (3, 4, 5))
        day_num = start_day + offset
        offset += 1
        if not any((disp, daily, combined, nim, total)):
            continue
        actual_year, actual_month = year, month
        while day_num > calendar.monthrange(actual_year, actual_month)[1]:
            day_num -= calendar.monthrange(actual_year, actual_month)[1]
            actual_year, actual_month = ((actual_year + 1, 1) if actual_month == 12
                                         else (actual_year, actual_month + 1))
        day = dt.date(actual_year, actual_month, day_num)
        if not combined:
            combined = disp + daily
        if not total:
            total = combined + nim
        yield day, disp, daily, nim, total, combined, None


def install(module):
    """Switch served routes once; do not modify the standalone calculator."""
    if getattr(module, '_daily_only_installed', False):
        return
    if not getattr(module, '_daily_monthly_ledger_installed', False):
        install_derived_months(module)

    def commit_daily_only(conn, user_id, year, month, commit=True):
        # app.input_sales historically calls recalc_monthly_summary to commit
        # its daily upsert. Preserve that transaction boundary, but never write
        # a redundant monthly row or a prior-month delta.
        if commit:
            conn.commit()

    module.recalc_monthly_summary = commit_daily_only

    old_upload = module.app.view_functions['upload_excel']

    @wraps(old_upload)
    def upload_daily_excel():
        if 'user_id' not in session:
            return redirect(url_for('login'))
        upload = request.files.get('excel_file')
        if not upload or not upload.filename:
            flash('파일을 선택해 주세요.', 'warning')
            return redirect(url_for('dashboard'))
        workbook = None
        conn = None
        try:
            data = upload.read(MAX_WORKBOOK_SIZE + 1)
            if not data or len(data) > MAX_WORKBOOK_SIZE:
                raise ValueError('파일은 10MB 이하의 엑셀만 업로드할 수 있습니다.')
            workbook = _open_workbook(data, request.form.get('excel_password', '').strip())
            conn = module.get_db()
            imported = 0
            for day, disp, daily, nim, total, combined, memo in _parse_daily_rows(
                    workbook, upload.filename, request.form.get('start_year')):
                _daily_upsert(conn, session['user_id'], day, disp, daily, nim,
                              total, dpd=combined, memo=memo)
                imported += 1
            if not imported:
                conn.rollback()
                flash('일장부에 가져올 날짜별 기록이 없습니다.', 'warning')
                return redirect(url_for('dashboard'))
            conn.commit()
        except Exception as exc:
            if conn is not None:
                conn.rollback()
            flash(f'일장부 가져오기 실패: {exc}', 'danger')
            return redirect(url_for('dashboard'))
        finally:
            if conn is not None:
                conn.close()
            if workbook is not None:
                workbook.close()
        module.invalidate_user_cache(session['user_id'])
        flash(f'일장부 {imported}일치 가져오기 완료. 월별 합계는 자동 계산됩니다.', 'success')
        return redirect(url_for('dashboard'))

    module.app.view_functions['upload_excel'] = upload_daily_excel

    original_forecast = module.get_month_forecast
    original_diagnosis = module.get_profit_balance_diagnosis

    def derived_row(conn, user_id, year, month):
        return next((r for r in module.get_cached_monthly_summary(conn, user_id)
                     if int(r['year']) == year and int(r['month']) == month), {})

    @wraps(original_forecast)
    def daily_forecast(conn, user_id, year, month, *args, **kwargs):
        # Passing an explicit empty dict suppresses the old function's
        # monthly_summary fallback SQL even when there are no daily entries.
        kwargs['month_summary_row'] = derived_row(conn, user_id, year, month)
        if kwargs.get('last_year_total') is None:
            prior = derived_row(conn, user_id, year - 1, month)
            kwargs['last_year_total'] = int(prior.get('grand_total') or 0)
        return original_forecast(conn, user_id, year, month, *args, **kwargs)

    @wraps(original_diagnosis)
    def daily_diagnosis(conn, user_id, year, month, *args, **kwargs):
        kwargs['month_summary_row'] = derived_row(conn, user_id, year, month)
        return original_diagnosis(conn, user_id, year, month, *args, **kwargs)

    module.get_month_forecast = daily_forecast
    module.get_profit_balance_diagnosis = daily_diagnosis

    # These old helpers used to SELECT monthly_summary directly. Replace their
    # served-app entrypoints as well so no retired monthly amount leaks back.
    def latest_from_daily(user_id, conn=None):
        rows = module.get_cached_monthly_summary(conn, user_id)
        found = next((r for r in reversed(rows) if int(r['grand_total']) > 0), None)
        if found is None:
            today = dt.date.today()
            return {'year':today.year, 'month':today.month, 'dispensing_plus_daily':0,
                    'non_insurance':0, 'extra_profit_total':0, 'grand_total':0, 'diff':0}
        return {'year':found['year'], 'month':found['month'],
                'dispensing_plus_daily':found['dispensing_plus_daily_total'],
                'non_insurance':found['non_insurance_total'],
                'extra_profit_total':found['extra_profit_total'],
                'grand_total':found['grand_total'],
                'diff':found['prev_month_diff'] or 0}

    module.get_latest_month_summary = latest_from_daily
    module._daily_only_installed = True
