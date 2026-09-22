import datetime as dt
import importlib
import sqlite3
import sys

import pytest


@pytest.fixture
def service(tmp_path, monkeypatch):
    monkeypatch.delenv('DATABASE_URL', raising=False)
    import database
    monkeypatch.setattr(database, 'SQLITE_PATH', str(tmp_path / 'sales.db'))
    if 'app' not in sys.modules:
        module = importlib.import_module('app')
    else:
        module = sys.modules['app']
        database.init_db()
    import profit_analysis
    monkeypatch.setattr(profit_analysis, 'korea_today', lambda: dt.date(2026, 9, 19))
    monkeypatch.setattr(module, 'korea_today', lambda: dt.date(2026, 9, 19))
    module._USER_CACHE.clear()
    module.app.config.update(TESTING=True)
    return module


@pytest.fixture
def conn(service):
    connection = service.get_db()
    yield connection
    connection.close()


def add(conn, day, disp=100, daily=20, nim=10, user=1):
    date = dt.date.fromisoformat(day)
    conn.execute('''INSERT INTO daily_profit
        (user_id,date,day_of_week,dispensing_fee,daily_net_profit,
        dispensing_plus_daily,non_insurance_margin,total)
        VALUES (?,?,?,?,?,?,?,?)''',
        (user, day, '월화수목금토일'[date.weekday()], disp, daily, disp+daily, nim, disp+daily+nim))
    conn.commit()


def test_past_gaps_are_not_remaining_business_days(service, conn):
    add(conn, '2026-09-01')
    result = service.get_month_forecast(conn, 1, 2026, 9, as_of=dt.date(2026, 9, 29))
    assert result['remaining_business_days'] == 2
    assert result['missing_count'] == 23
    assert result['yoy_growth_pct'] is None


def test_only_recent_eight_weeks_and_zero_days_enter_average(service, conn):
    add(conn, '2026-06-01', 100000, 0, 0)  # Monday, outside window
    add(conn, '2026-08-24', 100, 0, 0)
    add(conn, '2026-08-31', 0, 0, 0)
    result = service.get_month_forecast(conn, 1, 2026, 9, as_of=dt.date(2026, 9, 28))
    assert result['weekday_samples']['월']['mean'] == 50
    assert result['weekday_samples']['월']['count'] == 2
    assert result['weekday_samples']['월']['low'] == 0


def test_schedule_includes_sunday_and_excludes_explicit_holiday(service, conn):
    schedule = {'weekdays': [0,1,2,3,4,5,6], 'closed_dates': ['2026-09-28'], 'open_dates': []}
    result = service.get_month_forecast(conn, 1, 2026, 9, as_of=dt.date(2026, 9, 27), schedule=schedule)
    assert result['remaining_business_days'] == 3  # Sun, Tue, Wed
    assert result['forecast_total'] is None  # no observations: not a zero forecast


def test_past_month_is_never_extrapolated(service, conn):
    add(conn, '2026-08-03')
    result = service.get_month_forecast(conn, 1, 2026, 8, as_of=dt.date(2026, 9, 19))
    assert result['remaining_business_days'] == 0
    assert result['expected_additional'] == 0
    assert result['forecast_total'] == 130
    assert result['is_past'] is True


def test_forecast_excludes_future_entries_and_other_users(service, conn):
    add(conn, '2026-09-28', 999999, 0, 0)
    result = service.get_month_forecast(conn, 1, 2026, 9, as_of=dt.date(2026, 9, 19))
    assert result['current_total'] == 0
    assert result['future_entry_count'] == 1
    assert result['forecast_total'] is None


def test_monthly_only_balance_does_not_invent_breakdown(service, conn):
    summary = {'grand_total': 1000, 'dispensing_plus_daily_total': 900, 'non_insurance_total': 100}
    result = service.get_profit_balance_diagnosis(conn, 1, 2026, 9, entered_rows=[], month_summary_row=summary)
    assert result['has_breakdown'] is False
    assert result['disp_val'] is None
    assert result['daily_val'] is None
    assert result['total'] == 1000


def test_briefing_reports_decline_and_day_adjusted_change(service, conn):
    add(conn, '2026-08-01', 100, 0, 0)
    add(conn, '2026-08-02', 100, 0, 0)
    add(conn, '2026-09-01', 80, 0, 0)
    add(conn, '2026-09-02', 70, 0, 0)
    result = service.get_ai_narrative_briefing(conn, 1, {}, None)
    comparison = result['month_comparison']
    assert comparison['diff'] == -50
    assert comparison['daily_avg_diff'] == -25
    assert comparison['components'][0]['diff'] == -50
    assert '감소' in ' '.join(result['paragraphs'])
    assert '방어선' not in ' '.join(result['paragraphs'])


def test_unequal_input_days_are_explicit_not_called_business_days(service, conn):
    add(conn, '2026-08-01', 100, 0, 0)
    add(conn, '2026-08-02', 100, 0, 0)
    add(conn, '2026-09-02', 100, 0, 0)
    result = service.get_ai_narrative_briefing(conn, 1, {}, None)['month_comparison']
    assert result['current_days'] == 1
    assert result['previous_days'] == 2
    assert result['daily_avg_diff'] == 0
    assert result['diff'] == -100


def test_year_comparison_missing_does_not_claim_zero_growth(service, conn):
    add(conn, '2026-09-01')
    result = service.get_yoy_day_comparison(conn, 1, '2026-09-01')
    assert result['growth_pct'] is None
    assert result['yoy_total'] is None


def test_annual_briefing_does_not_claim_stability_from_decline(service):
    rows = [{'month': 1, 'grand_total': 100}]
    result = service.generate_annual_narrative_report(2026, rows,
        {'grand':100,'dpd':90,'nim':10}, {'grand':1000}, rows[0], rows[0], [])
    text = ' '.join(result)
    assert '감소' in text
    assert '방어선' not in text
    assert '비교 기간' in text


def client_for(service, user=1):
    client = service.app.test_client()
    with client.session_transaction() as session:
        session.update(user_id=user, username='tester', pharmacy_name='Test', schedule_csrf='test-token')
    return client


def test_schedule_persists_and_is_user_scoped(service, conn):
    client = client_for(service)
    response = client.post('/business-schedule', data={
        'csrf_token':'test-token','weekdays':['0','6'], 'closed_dates':'2026-09-28',
        'open_dates':'2026-09-29', 'year':'2026','month':'9'})
    assert response.status_code == 302
    from profit_analysis import get_business_schedule
    assert get_business_schedule(conn, 1)['weekdays'] == [0,6]
    assert get_business_schedule(conn, 2)['weekdays'] == [0,1,2,3,4,5]
    result = service.get_month_forecast(conn, 1, 2026, 9, as_of=dt.date(2026,9,27))
    assert result['remaining_business_days'] == 2


def test_invalid_schedule_is_rejected_without_overwrite(service, conn):
    client = client_for(service)
    for data in [
        {'weekdays':['8']}, {'closed_dates':'2026-02-30'},
        {'closed_dates':'2026-09-28','open_dates':'2026-09-28'},
    ]:
        response = client.post('/business-schedule', data={'csrf_token':'test-token',**data})
        assert response.status_code == 400
    assert conn.execute('SELECT COUNT(*) FROM business_schedules').fetchone()[0] == 0
    assert client.post('/business-schedule', data={'weekdays':['0']}).status_code == 400


def test_pages_render_no_data_and_summary_only(service, conn):
    client = client_for(service)
    for url in ['/', '/calendar?year=2026&month=9', '/report', '/calculator']:
        response = client.get(url)
        assert response.status_code == 200
    conn.execute('INSERT INTO monthly_summary (user_id,year,month,grand_total,dispensing_plus_daily_total,non_insurance_total) VALUES (1,2026,9,1000,900,100)')
    conn.commit()
    service._USER_CACHE.clear()
    response = client.get('/')
    assert response.status_code == 200
    assert '세부 자료 없음' in response.get_data(as_text=True)


def test_pages_render_actuals_and_missing_last_year(service, conn):
    add(conn,'2026-09-01')
    service.recalc_monthly_summary(conn,1,2026,9)
    client = client_for(service)
    for url in ['/', '/calendar?year=2026&month=9', '/input', '/report']:
        response = client.get(url)
        assert response.status_code == 200


def test_complete_month_forecast_has_hand_checked_range(service, conn):
    for day, amount in [('2026-09-07', 100), ('2026-09-14', 200), ('2026-09-21', 300)]:
        add(conn, day, amount, 0, 0)
    schedule = {'weekdays':[0], 'closed_dates':[], 'open_dates':[]}
    result = service.get_month_forecast(conn,1,2026,9,as_of=dt.date(2026,9,28),schedule=schedule,last_year_total=400)
    assert result['missing_count'] == 0
    assert result['forecast_total'] == 800
    assert result['forecast_low'] == 700
    assert result['forecast_high'] == 900
    assert result['yoy_growth_pct'] == 100.0


def test_tenant_history_does_not_leak_into_forecast(service, conn):
    conn.execute("INSERT INTO users (id,username,password_hash,pharmacy_name) VALUES (2,'other','unused','Other')")
    add(conn,'2026-09-14',999999,0,0,user=2)
    result = service.get_month_forecast(conn,1,2026,9,as_of=dt.date(2026,9,28))
    assert result['sample_count'] == 0
    assert result['forecast_total'] is None


def test_month_summary_only_is_not_displayed_as_zero(service, conn):
    conn.execute('INSERT INTO monthly_summary (user_id,year,month,grand_total) VALUES (1,2026,8,700)')
    conn.commit()
    result = service.get_month_forecast(conn,1,2026,8,as_of=dt.date(2026,9,19))
    assert result['current_total'] == 700
    assert result['summary_only'] is True
    assert result['forecast_total'] == 700


def test_summary_only_does_not_claim_daily_amounts_are_excluded(service, conn):
    conn.execute('INSERT INTO monthly_summary (user_id,year,month,grand_total) VALUES (1,2026,8,700)')
    conn.commit()
    result = service.get_month_forecast(conn,1,2026,8,as_of=dt.date(2026,9,19),last_year_total=500)
    assert result['missing_count'] == 0
    assert result['yoy_growth_pct'] is None  # completeness is unknown
    client = client_for(service)
    text = client.get('/calendar?year=2026&month=8').get_data(as_text=True)
    assert '합계와 예측에 보충하지 않았습니다' not in text
    assert '일별 세부 자료 없음' in text


def test_dashboard_query_budget_and_fresh_values(service, conn, monkeypatch):
    for day in ('2025-09-20', '2026-08-19', '2026-09-12', '2026-09-19'):
        add(conn, day)
    service.recalc_monthly_summary(conn, 1, 2026, 9)
    statements = []
    original = service.get_db
    def tracked():
        c = original()
        c.set_trace_callback(lambda sql: statements.append(sql) if sql.lstrip().upper().startswith('SELECT') else None)
        return c
    monkeypatch.setattr(service, 'get_db', tracked)
    client = client_for(service)
    response = client.get('/')
    assert response.status_code == 200
    base_queries = [sql for sql in statements if 'extra_profit' not in sql]
    misc_queries = [sql for sql in statements if 'extra_profit' in sql]
    assert len(base_queries) <= 3, statements
    assert len(misc_queries) <= 2, statements
    statements.clear()
    conn.execute("UPDATE daily_profit SET total=99999 WHERE user_id=1 AND date='2026-09-19'")
    conn.commit()
    # Mirror every application write: commit, then invalidate the user's caches.
    service.invalidate_user_cache(1)
    response = client.get('/')
    assert '99,999' in response.get_data(as_text=True)
    base_queries = [sql for sql in statements if 'extra_profit' not in sql]
    misc_queries = [sql for sql in statements if 'extra_profit' in sql]
    assert len(base_queries) <= 3, statements
    assert len(misc_queries) <= 2, statements


def test_login_needs_no_external_render_dependencies(service):
    html = service.app.test_client().get('/login').get_data(as_text=True)
    assert 'https://cdn' not in html
    assert 'autocomplete="current-password"' in html


def test_non_chart_page_does_not_download_chart_library(service):
    html = client_for(service).get('/input').get_data(as_text=True)
    assert 'chart.umd.min.js' not in html


def test_snapshot_forecast_matches_direct_query_for_past_and_future(service, conn):
    for day in ('2025-01-02', '2025-02-03', '2026-08-31', '2026-09-19', '2026-12-01'):
        add(conn, day)
    for year, month in ((2025, 2), (2026, 9), (2026, 12)):
        snapshot = service.load_analysis_rows(conn, 1, year, month)
        assert service.get_month_forecast(conn, 1, year, month, history_rows=snapshot) == service.get_month_forecast(conn, 1, year, month)


def test_calendar_uses_one_connection(service, conn, monkeypatch):
    add(conn, '2026-09-19')
    service.recalc_monthly_summary(conn, 1, 2026, 9)
    calls = []
    original = service.get_db
    def tracked():
        calls.append(1)
        return original()
    monkeypatch.setattr(service, 'get_db', tracked)
    import profit_analysis
    monkeypatch.setattr(profit_analysis, 'get_db', tracked)
    assert client_for(service).get('/calendar?year=2026&month=9').status_code == 200
    assert len(calls) == 1


def test_login_success_and_failure_with_synthetic_credentials(service, conn):
    from werkzeug.security import generate_password_hash
    conn.execute('UPDATE users SET username=?, password_hash=? WHERE id=1',
                 ('speed-test-user', generate_password_hash('synthetic-test-password')))
    conn.commit()
    client = service.app.test_client()
    rejected = client.post('/login', data={'username': 'speed-test-user', 'password': 'wrong'})
    assert rejected.status_code == 200
    with client.session_transaction() as session:
        assert 'user_id' not in session
    accepted = client.post('/login', data={'username': 'speed-test-user', 'password': 'synthetic-test-password'})
    assert accepted.status_code == 302
    with client.session_transaction() as session:
        assert session['user_id'] == 1
    assert client.get('/').status_code == 200


def test_daily_save_rolls_back_when_month_summary_fails(service, conn):
    add(conn, '2026-09-19', disp=100)
    service.recalc_monthly_summary(conn, 1, 2026, 9)
    conn.execute("""CREATE TRIGGER fail_summary BEFORE INSERT ON monthly_summary
        BEGIN SELECT RAISE(ABORT, 'injected summary failure'); END""")
    conn.commit()
    response = client_for(service).post('/input', data={
        'date': '2026-09-19', 'dispensing_fee': '999',
        'daily_net_profit': '20', 'non_insurance_margin': '10'})
    assert response.status_code == 302
    assert conn.execute("SELECT total FROM daily_profit WHERE date='2026-09-19'").fetchone()['total'] == 130
    assert conn.execute('SELECT grand_total FROM monthly_summary').fetchone()['grand_total'] == 130


def excel_bytes(rows):
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(['날짜', '요일', '조제료', '일매순익', '비보험', '메모'])
    for row in rows:
        wb.active.append(row)
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    output.seek(0)
    return output


def test_excel_failure_rolls_back_all_months_and_returns_page(service, conn):
    conn.execute("""CREATE TRIGGER fail_second_month BEFORE INSERT ON monthly_summary
        WHEN (SELECT COUNT(*) FROM monthly_summary) > 0
        BEGIN SELECT RAISE(ABORT, 'injected second month failure'); END""")
    conn.commit()
    response = client_for(service).post('/upload_excel', data={'excel_file': (
        excel_bytes([['2026-08-01', '토', 100, 20, 10], ['2026-09-01', '화', 200, 30, 10]]), 'test.xlsx')})
    assert conn.execute('SELECT COUNT(*) FROM daily_profit').fetchone()[0] == 0
    assert conn.execute('SELECT COUNT(*) FROM monthly_summary').fetchone()[0] == 0
    assert response.status_code == 302


def test_invalid_excel_returns_error_message_instead_of_500(service):
    import io
    response = client_for(service).post('/upload_excel', data={
        'excel_file': (io.BytesIO(b'not an excel file'), 'broken.xlsx')}, follow_redirects=True)
    assert response.status_code == 200
    assert '엑셀 가져오기 실패' in response.get_data(as_text=True)


def test_edit_previous_month_refreshes_next_month_difference(service, conn):
    add(conn, '2026-08-01', disp=100, daily=0, nim=0)
    add(conn, '2026-09-01', disp=300, daily=0, nim=0)
    service.recalc_monthly_summary(conn, 1, 2026, 8)
    service.recalc_monthly_summary(conn, 1, 2026, 9)
    response = client_for(service).post('/input', data={'date':'2026-08-01', 'dispensing_fee':'200'})
    assert response.status_code == 302
    row = conn.execute('SELECT prev_month_diff FROM monthly_summary WHERE month=9').fetchone()
    assert row['prev_month_diff'] == 100


def test_diagnostics_require_admin(service):
    assert service.app.test_client().get('/debug/perf').status_code == 302
    assert client_for(service).get('/debug/perf').status_code == 302


def test_login_startup_does_not_import_excel_dependencies(tmp_path):
    import os
    import subprocess
    code = '''
import sys, importlib.abc
class BlockExcel(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in ('openpyxl', 'msoffcrypto'):
            raise ImportError('Excel must not load during login startup')
sys.meta_path.insert(0, BlockExcel())
import database
database.SQLITE_PATH = sys.argv[1]
import app
assert app.app.test_client().get('/login').status_code == 200
'''
    env = dict(os.environ)
    env.pop('DATABASE_URL', None)
    result = subprocess.run([sys.executable, '-c', code, str(tmp_path / 'startup.db')],
                            capture_output=True, text=True, env=env)
    assert result.returncode == 0, result.stderr


def test_diagnostics_return_connection_on_query_failure(service, conn, monkeypatch):
    conn.set_authorizer(lambda *args: sqlite3.SQLITE_DENY)
    monkeypatch.setattr(service, 'get_db', lambda: conn)
    with service.app.test_request_context('/debug/perf'):
        service.session.update(user_id=1, username='admin')
        with pytest.raises(sqlite3.DatabaseError):
            service.debug_perf()
    with pytest.raises(sqlite3.ProgrammingError, match='closed'):
        conn.execute('SELECT 1')


def test_standard_excel_commits_once_and_refreshes_cached_totals(service, conn, monkeypatch):
    statements = []
    original = service.get_db
    def tracked():
        c = original()
        c.set_trace_callback(statements.append)
        return c
    monkeypatch.setattr(service, 'get_db', tracked)
    assert service.get_cached_monthly_summary(conn, 1) == []
    response = client_for(service).post('/upload_excel', data={'excel_file': (
        excel_bytes([['2026-09-01', '화', 200, 30, 10], ['2026-08-01', '토', 100, 20, 10]]), 'test.xlsx')})
    assert response.status_code == 302
    assert sum(s.strip().upper() == 'COMMIT' for s in statements) == 1
    rows = service.get_cached_monthly_summary(conn, 1)
    assert [r['grand_total'] for r in rows] == [130, 240]
    assert rows[1]['prev_month_diff'] == 110
    assert conn.execute('SELECT COUNT(*) FROM daily_profit').fetchone()[0] == 2


def test_excel_downloads_still_open_after_lazy_import(service):
    import io
    import openpyxl
    client = client_for(service)
    with client.session_transaction() as session:
        session['username'] = 'admin'
    for url in ('/download_template', '/admin/export_backup'):
        response = client.get(url)
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.data))
        assert workbook.active.max_row >= 1
        workbook.close()


def test_lost_connection_during_save_preserves_failure_response(service, monkeypatch):
    import database
    class Disconnected:
        def execute(self, *args):
            raise database.psycopg2.OperationalError('connection lost during save')
        def rollback(self):
            raise database.psycopg2.InterfaceError('connection already closed')
        def close(self):
            pass
    monkeypatch.setattr(service, 'get_db', Disconnected)
    response = client_for(service).post('/input', data={
        'date': '2026-09-19', 'dispensing_fee': '100'})
    assert response.status_code == 302
