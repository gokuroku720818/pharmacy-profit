"""Synthetic-only checks: independent monthly ledger must not be read or written."""
import datetime as dt
import io

import openpyxl

from test_analysis import add, client_for, conn, service


def _install(service, monkeypatch):
    import profit_display
    import profit_analysis
    for name in ('get_cached_monthly_summary', 'recalc_monthly_summary',
                 'get_month_forecast', 'get_profit_balance_diagnosis'):
        monkeypatch.setattr(service, name, getattr(service, name))
    monkeypatch.setattr(profit_analysis, 'get_month_forecast', profit_analysis.get_month_forecast)
    monkeypatch.setattr(profit_analysis, 'get_profit_balance_diagnosis', profit_analysis.get_profit_balance_diagnosis)
    for name in ('export_csv', 'historical_reconciliation', 'upload_excel'):
        monkeypatch.setitem(service.app.view_functions, name, service.app.view_functions[name])
    monkeypatch.setattr(profit_display, 'flash', profit_display.flash)
    monkeypatch.setattr(profit_display, '_daily_warning_installed', False, raising=False)
    monkeypatch.setattr(service, '_daily_monthly_ledger_installed', False, raising=False)
    monkeypatch.setattr(service, '_daily_only_installed', False, raising=False)
    from daily_only import install
    install(service)
    return client_for(service)


def _legacy(conn):
    conn.execute('''INSERT INTO monthly_summary
        (user_id,year,month,dispensing_plus_daily_total,non_insurance_total,grand_total,prev_month_diff)
        VALUES (1,2026,8,1,1,999999,0)''')
    conn.commit()


def test_daily_input_writes_only_daily_record_not_retired_monthly_table(service, conn, monkeypatch):
    _legacy(conn)
    client = _install(service, monkeypatch)
    response = client.post('/input', data={'date':'2026-08-01','dispensing_fee':'400',
                        'daily_net_profit':'100','non_insurance_margin':'50'})
    assert response.status_code == 302
    assert conn.execute("SELECT total FROM daily_profit WHERE user_id=1 AND date='2026-08-01'").fetchone()[0] == 550
    assert conn.execute('SELECT grand_total FROM monthly_summary WHERE user_id=1 AND year=2026 AND month=8').fetchone()[0] == 999999
    assert client.get('/report?year=2026').status_code == 200
    assert service.get_cached_monthly_summary(None,1)[0]['grand_total'] == 550


def test_legacy_workbook_import_ignores_summary_sheet_and_only_writes_daily(service, conn, monkeypatch):
    _legacy(conn)
    client = _install(service, monkeypatch)
    wb = openpyxl.Workbook()
    weekly = wb.active
    weekly.title = '주표'
    weekly.append(['8월 1일~'])
    weekly.append(['토', 400, 100, 500, 50, 550])
    monthly = wb.create_sheet('합계')
    monthly.append(['연도','월','조제+일매','비보험','월장부 총액','차이'])
    monthly.append([2026,'8월',1000,2000,99999999,3000])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    response = client.post('/upload_excel', data={'excel_file':(out,'weekly_2026.xlsx')},
                           content_type='multipart/form-data')
    assert response.status_code == 302
    assert conn.execute("SELECT total FROM daily_profit WHERE user_id=1 AND date='2026-08-01'").fetchone()[0] == 550
    assert conn.execute('SELECT grand_total FROM monthly_summary WHERE user_id=1 AND month=8').fetchone()[0] == 999999
    assert service.get_cached_monthly_summary(None,1)[0]['grand_total'] == 550


def test_standard_workbook_only_daily_and_missing_year_rejected_without_guess(service, conn, monkeypatch):
    client = _install(service, monkeypatch)
    standard = openpyxl.Workbook()
    ws = standard.active
    ws.append(['날짜','요일','조제료','일매순익','비보험','메모'])
    ws.append(['2026-09-01','화',100,20,10,'daily'])
    buf = io.BytesIO()
    standard.save(buf)
    buf.seek(0)
    assert client.post('/upload_excel',data={'excel_file':(buf,'daily.xlsx')},
                       content_type='multipart/form-data').status_code == 302
    assert conn.execute("SELECT total FROM daily_profit WHERE user_id=1 AND date='2026-09-01'").fetchone()[0] == 130
    assert conn.execute('SELECT COUNT(*) FROM monthly_summary').fetchone()[0] == 0
    legacy = openpyxl.Workbook()
    legacy.active.append(['8월 1일~'])
    legacy.active.append(['토',100,20,120,10,130])
    buf = io.BytesIO()
    legacy.save(buf)
    buf.seek(0)
    client.post('/upload_excel',data={'excel_file':(buf,'unknown.xlsx')},content_type='multipart/form-data')
    assert conn.execute('SELECT COUNT(*) FROM daily_profit').fetchone()[0] == 1


def test_forecast_and_diagnosis_never_read_retired_monthly_only_amount(service, conn, monkeypatch):
    _legacy(conn)
    client = _install(service, monkeypatch)
    forecast = service.get_month_forecast(conn,1,2026,8,as_of=dt.date(2026,9,19))
    assert forecast['current_total'] == 0
    assert forecast['summary_only'] is False
    assert forecast['forecast_total'] == 0
    diagnosis = service.get_profit_balance_diagnosis(conn,1,2026,8,entered_rows=[])
    assert diagnosis['total'] == 0
    assert client.get('/report?year=2026').status_code == 200
    assert '999,999' not in client.get('/report?year=2026').get_data(as_text=True)


def test_daily_only_install_is_idempotent_and_keeps_calculator_unmodified(service, monkeypatch):
    from daily_only import install
    before_calculator = service.app.view_functions['calculator']
    client = _install(service, monkeypatch)
    before_daily_upload = service.app.view_functions['upload_excel']
    install(service)
    assert service.app.view_functions['upload_excel'] is before_daily_upload
    assert service.app.view_functions['calculator'] is before_calculator
    assert client.get('/calculator').status_code == 200
