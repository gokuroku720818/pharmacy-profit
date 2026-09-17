"""
엑셀 순익표 데이터를 SQLite DB로 마이그레이션하는 스크립트

엑셀 구조:
- 시트1 (주표): 주간별 요일(월~토) 순익 데이터
  - B열: 조제료
  - C열: 일매순익
  - D열: 조제 + 일매순익
  - E열: 비보험약가차액순익
  - F열: 전체 합계
  - 주간 구분 행: "10월 17~22 <3주차>" 형태
  - 합계 행: 주간 평균/합계 (F열에 주합계)

- 시트2 (합계): 월별 합산
  - A열: 연도, B열: 월
  - C열: 조제+일매순익 월합계
  - D열: 비보험약가차액 월합계
  - E열: 전체 합계
  - F열: 전월 대비 증감

- 시트3 (Sheet1): 일별 합계를 월 단위 가로배열
"""
import msoffcrypto
import io
import re
import datetime
import calendar
import openpyxl
import os
from database import get_db, init_db


def decrypt_excel(filepath, password):
    """암호화된 엑셀 파일을 복호화"""
    with open(filepath, 'rb') as f:
        file = msoffcrypto.OfficeFile(f)
        file.load_key(password=password)
        decrypted = io.BytesIO()
        file.decrypt(decrypted)
        decrypted.seek(0)
    return decrypted


def safe_int(val):
    """안전하게 정수 변환"""
    if val is None:
        return 0
    try:
        return int(val)
    except (ValueError, TypeError):
        return 0


def import_monthly_summary(wb):
    """합계 시트에서 월별 데이터 임포트"""
    ws = wb[wb.sheetnames[1]]  # 두 번째 시트 = 합계
    conn = get_db()
    cursor = conn.cursor()

    imported = 0
    for r in range(2, ws.max_row + 1):
        year = ws.cell(r, 1).value
        month_str = str(ws.cell(r, 2).value or '')

        if year is None or not isinstance(year, (int, float)):
            continue

        year = int(year)
        month_match = re.search(r'(\d+)', month_str)
        if not month_match:
            continue
        month = int(month_match.group(1))

        dpd = safe_int(ws.cell(r, 3).value)    # 조제+일매순익 합계
        nim = safe_int(ws.cell(r, 4).value)     # 비보험약가차액 합계
        grand = safe_int(ws.cell(r, 5).value)   # 전체 합계
        diff = safe_int(ws.cell(r, 6).value)    # 전월 대비

        if grand == 0 and dpd == 0 and nim == 0:
            continue

        cursor.execute('''
            INSERT OR REPLACE INTO monthly_summary
            (year, month, dispensing_plus_daily_total, non_insurance_total, grand_total, prev_month_diff)
            VALUES (?, ?, ?, ?, ?, ?)
        ''', (year, month, dpd, nim, grand, diff))
        imported += 1

    conn.commit()
    conn.close()
    print(f"월별 합계 {imported}건 임포트 완료")


def import_daily_from_weekly(wb):
    """주표 시트에서 일별 데이터 임포트"""
    ws = wb[wb.sheetnames[0]]
    conn = get_db()
    cursor = conn.cursor()

    DAY_NAMES = ['월', '화', '수', '목', '금', '토']

    imported = 0
    current_year = 2022
    current_month = 10
    current_start_day = 17
    day_offset = 0

    for r in range(1, ws.max_row + 1):
        cell_a = ws.cell(r, 1).value

        if cell_a is None:
            day_offset = 0
            continue

        cell_a_str = str(cell_a).strip()

        # 주간 헤더 감지: '10월 17~22', '1월2일 ~ 1월7일', '1월 5 ~ 1월 10일' 등 모두 매칭
        week_match = re.search(r'(\d+)\s*월?\s*(\d+)\s*일?\s*~', cell_a_str)
        if week_match:
            new_month = int(week_match.group(1))
            current_start_day = int(week_match.group(2))
            day_offset = 0

            # 11~12월에서 1~2월로 넘어갈 때 연도 증가
            if new_month < current_month and new_month <= 2 and current_month >= 11:
                current_year += 1
            current_month = new_month
            continue

        # 요일 행 감지
        if cell_a_str in DAY_NAMES:
            dispensing = safe_int(ws.cell(r, 2).value)      # 조제료
            daily_net = safe_int(ws.cell(r, 3).value)       # 일매순익
            dpd = safe_int(ws.cell(r, 4).value)             # 조제+일매순익
            nim = safe_int(ws.cell(r, 5).value)             # 비보험약가차액순익
            total = safe_int(ws.cell(r, 6).value)           # 전체 합계

            # 빈 템플릿 행(모두 0원)은 제외
            if dispensing == 0 and daily_net == 0 and nim == 0 and total == 0:
                day_offset += 1
                continue

            if dpd == 0 and (dispensing > 0 or daily_net > 0):
                dpd = dispensing + daily_net
            if total == 0:
                total = dpd + nim

            day = current_start_day + day_offset
            day_offset += 1

            actual_month = current_month
            actual_year = current_year
            max_day = calendar.monthrange(actual_year, actual_month)[1]
            if day > max_day:
                day = day - max_day
                actual_month += 1
                if actual_month > 12:
                    actual_month = 1
                    actual_year += 1

            try:
                dt = datetime.date(actual_year, actual_month, day)
                date_str = dt.isoformat()

                cursor.execute('''
                    INSERT OR REPLACE INTO daily_profit
                    (date, day_of_week, dispensing_fee, daily_net_profit,
                     dispensing_plus_daily, non_insurance_margin, total, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, datetime('now', 'localtime'))
                ''', (date_str, cell_a_str, dispensing, daily_net, dpd, nim, total))
                imported += 1
            except ValueError:
                pass

    conn.commit()
    conn.close()
    print(f"일별 순익 {imported}건 임포트 완료")


def main():
    excel_path = os.path.join(os.path.expanduser('~'), 'Desktop', '순익표_최종.xlsx')
    password = '7581'

    if not os.path.exists(excel_path):
        print(f"파일을 찾을 수 없습니다: {excel_path}")
        return

    print("1. 데이터베이스 초기화...")
    init_db()

    print("2. 엑셀 파일 복호화...")
    decrypted = decrypt_excel(excel_path, password)
    wb = openpyxl.load_workbook(decrypted, data_only=True)
    print(f"   시트: {wb.sheetnames}")

    print("3. 월별 합계 데이터 임포트...")
    import_monthly_summary(wb)

    print("4. 일별 순익 데이터 임포트...")
    import_daily_from_weekly(wb)

    # 결과 확인
    conn = get_db()
    ms_count = conn.execute("SELECT COUNT(*) FROM monthly_summary").fetchone()[0]
    ds_count = conn.execute("SELECT COUNT(*) FROM daily_profit").fetchone()[0]
    total = conn.execute("SELECT SUM(grand_total) FROM monthly_summary").fetchone()[0]

    # 샘플 데이터 확인
    sample = conn.execute("SELECT * FROM daily_profit ORDER BY date LIMIT 5").fetchall()
    print("\n--- 샘플 데이터 ---")
    for s in sample:
        print(f"  {s['date']} ({s['day_of_week']}) 조제료:{s['dispensing_fee']:,} 일매순익:{s['daily_net_profit']:,} 비보험:{s['non_insurance_margin']:,} 합계:{s['total']:,}")

    conn.close()

    print(f"\n=== 임포트 완료 ===")
    print(f"월별 합계: {ms_count}건")
    print(f"일별 순익: {ds_count}건")
    print(f"누적 총순익: {total:,}원")


if __name__ == '__main__':
    main()
